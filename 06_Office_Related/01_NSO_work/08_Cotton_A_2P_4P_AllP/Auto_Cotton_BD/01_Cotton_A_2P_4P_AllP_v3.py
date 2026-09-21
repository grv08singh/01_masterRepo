from pathlib import Path
import numpy as np
import pandas as pd
import pyodbc
import sys
import os
import time
from datetime import date
import re
import openpyxl
from openpyxl import Workbook, load_workbook
from openpyxl.utils.dataframe import dataframe_to_rows
import warnings as wr
wr.filterwarnings('ignore')




def read_db():
    curr_path = Path.cwd()
    for file in curr_path.iterdir():
        if file.suffix == ".mdb" and 'NSSAS' in file.stem:
        #create connection with mdb database
            conn_str = (
                r'DRIVER={Microsoft Access Driver (*.mdb, *.accdb)};'
                f'DBQ={file};'
            )
            conn = pyodbc.connect(conn_str)
            
            ##get data from the table - EXPTMASTER
            cols_ml = ['EXPTID', 'SEASONCODE', 'SAMPLE', 
                       'STATE', 'ROCODE', 'SROCODE', 
                       'DISTRICT', 'DISTRICTNAME', 'SELORDER', 
                       'EXPT', 'CROPCODE', 'CROPNAME']
            sql_query_ml = 'SELECT * FROM "EXPTMASTER"'
            df_ml = pd.read_sql(sql_query_ml, conn)[cols_ml]
            
            #get data from the table - Sch20RC7
            cols7 = ['EXPTID', 'STAGIN', 'CROP', 'QTYP', 'STYPE', 'IRR', 
                     'FERTC', 'MANUR', 'PEST', 'CROPER', 'CROPCON', 'AFFECT']
    
            sql_query7 = 'SELECT * FROM "Sch20RC7"'
            df_rc7 = pd.read_sql(sql_query7, conn)[cols7]
    
            #get data from the table - Sch20Block5p2
            cols_5p2 = ['EXPTID', 'SRL', 'WEIGHT']
            sql_query_5p2 = 'SELECT * FROM "Sch20Block5p2"'
            df_5p2 = pd.read_sql(sql_query_5p2, conn)[cols_5p2]
            
            conn.close()

            return df_ml, df_rc7, df_5p2



def cotton_only(df_ml, df_rc7, df_5p2):
    crop_list = ['1211','1212','5511','5512','5611','5612','5711','5712']
    df_ml = df_ml[df_ml['CROPCODE'].isin(crop_list)]
    expt_id_list = list(set(list(df_ml['EXPTID'])))

    df_rc7 = df_rc7[df_rc7['CROP'].isin(['12','55','56','57'])]
    
    df_5p2 = df_5p2[df_5p2['EXPTID'].isin(expt_id_list)]

    df_ml_cat = pd.merge(df_ml, df_rc7, how='left', on='EXPTID')
    return df_ml_cat, df_5p2




def get_pickings(df_ml, df_5p2):
    expt_id_list = list(set(list(df_ml['EXPTID'])))
    for exp_id in expt_id_list:
        #getting df for just one single experiment
        df_5p2_single_expt = df_5p2[df_5p2['EXPTID'] == exp_id]

        #calculating 2-pickings sum
        pickings_2 = 0
        list_2 = [1,2]
        df_5p2_single_expt_2pickings = df_5p2_single_expt[df_5p2_single_expt['SRL'].isin(list_2)]
        pickings_2 = df_5p2_single_expt_2pickings['WEIGHT'].sum()
        
        #calculating 4-pickings sum
        pickings_4 = 0
        list_4 = [1,2,3,4]
        df_5p2_single_expt_4pickings = df_5p2_single_expt[df_5p2_single_expt['SRL'].isin(list_4)]
        pickings_4 = df_5p2_single_expt_4pickings['WEIGHT'].sum()

        df_ml.loc[(df_ml['EXPTID']==exp_id), '2_pickings'] = pickings_2*1000
        df_ml.loc[(df_ml['EXPTID']==exp_id), '4_pickings'] = pickings_4*1000

        #All picking sum
    df_ml['All_pickings'] = df_ml['QTYP']
    df_ml['HY'] = df_ml['STYPE']
    df_ml['IR'] = df_ml['IRR']
    df_ml['F'] = df_ml['FERTC']
    df_ml['M'] = df_ml['MANUR']
    df_ml['P'] = df_ml['PEST']
    df_ml['COND'] = df_ml['CROPCON']
    df_ml['AFF'] = df_ml['AFFECT']
    return df_ml





def save_xlsx(df):
            
    season_crop = {
        1: ['1211','1212','5511','5512'],
        4: ['5611','5612','5711','5712']
    }
    
    for season,crops in season_crop.items():
        if season==1:
            xl_name = "_KH_Cotton.xlsx"
        elif season==4:
            xl_name = "_RB_Cotton.xlsx"
        else:
            print('No Valid Season')
        
        #Delete cotton if exists from before
        if os.path.exists(xl_name):
            os.remove(xl_name)
        
        #Create cotton excel file
        wb = openpyxl.Workbook()
        
        for crop in crops:
            df_ml_crop = df_ml[df_ml['CROPCODE']==crop]

            sheet_title = crop
            ws1 = wb.create_sheet(title=sheet_title)
            if crop in ('1211','1212','5611','5612'):
                df_useful = df_ml_crop[['SAMPLE','DISTRICTNAME', 'SELORDER', 'EXPT',
                                       'CROPER', '2_pickings', 'All_pickings', 'HY',
                                       'IR', 'F', 'M', 'P', 'COND', 'AFF','STAGIN']]
            else:
                df_useful = df_ml_crop[['SAMPLE','DISTRICTNAME', 'SELORDER', 'EXPT',
                                       'CROPER', 'All_pickings', 'HY',
                                       'IR', 'F', 'M', 'P', 'COND', 'AFF','STAGIN']]
            
            df_useful.sort_values(['SAMPLE','DISTRICTNAME','SELORDER','EXPT'])
            obj_to_num = ['SELORDER','EXPT','HY','IR','F','M','P','COND','AFF']
            df_useful[obj_to_num] = df_useful[obj_to_num].apply(pd.to_numeric, errors='coerce')
            
            for r in dataframe_to_rows(df_useful, index=False, header=True):
                ws1.append(r)

        #remove unnecessary sheet
        sheet_to_remove = 'Sheet'
        if sheet_to_remove in wb.sheetnames:
            del wb[sheet_to_remove]
        
        wb.save(xl_name)




if __name__ == "__main__":
    #read database
    df_ml, df_rc7, df_5p2 = read_db()
    #filter cotton only
    df_ml, df_5p2 = cotton_only(df_ml, df_rc7, df_5p2)
    #get 2,4,All pickings
    df_ml = get_pickings(df_ml, df_5p2)
    #save excel files
    save_xlsx(df_ml)